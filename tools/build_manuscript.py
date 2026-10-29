"""Fill a private copy of the supplied template using frozen numerical outputs.

Run with the bundled document Python runtime (python-docx 1.2.0). This script
reads no raw recordings. The retained template is never modified.
"""
import argparse, csv, hashlib, json, platform, subprocess
from copy import deepcopy
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser=argparse.ArgumentParser()
parser.add_argument('--template',required=True)
parser.add_argument('--output-root',required=True)
parser.add_argument('--previous-report')
args=parser.parse_args()
root=Path(args.output_root).resolve();assets=root/'report_assets'
repo=Path(__file__).resolve().parents[1]
def rows(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
processing=json.loads((root/'processing_receipt.json').read_text())
evaluation=json.loads((root/'evaluation/evaluation_receipt.json').read_text())
summary=rows(root/'evaluation/metrics_summary.csv');results=rows(root/'evaluation/results.csv')
coverage=rows(root/'coverage.csv');sens=rows(root/'sensitivity_summary.csv');abl=rows(root/'ablation_summary.csv')
refs=json.loads((repo/'references/verified_references.json').read_text(encoding='utf-8'))
pooled=json.loads((assets/'pooled_counts.json').read_text())
paired=rows(root/'evaluation/paired_differences.csv')
if len(results)!=evaluation['outer_folds']*6:raise RuntimeError('incomplete model/fold table')
def metric(model,name):return next(r for r in summary if r['model']==model and r['metric']==name)
def ms(model,name):
    r=metric(model,name);return f"{float(r['mean']):.3f} ± {float(r['sd']):.3f}"
def ci(model,name):
    r=metric(model,name);return f"{float(r['ci95_lower']):.3f}–{float(r['ci95_upper']):.3f}"

template=Path(args.template).resolve();template_bytes=template.read_bytes();doc=Document(template)
original_headings={p.text.strip():deepcopy(p._p) for p in doc.paragraphs if p.style.name.startswith('Heading')}
table_patterns=[deepcopy(t._tbl) for t in doc.tables]
geometry=[{'width':s.page_width,'height':s.page_height,'top':s.top_margin,'bottom':s.bottom_margin,'left':s.left_margin,'right':s.right_margin} for s in doc.sections]
# All editable body slots are replaced. Section properties, styles, theme,
# header/footer parts, logos and relationship-linked page furniture survive.
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
manuscript=root/'manuscript';manuscript.mkdir(exist_ok=True)
content=[]
def p(text,style=None):
    paragraph=doc.add_paragraph(text,style=style);paragraph.paragraph_format.space_after=Pt(6)
    content.append(text);return paragraph
def h(text,level=1):
    if text in original_headings:
        node=deepcopy(original_headings[text]);body.insert(len(body)-1,node)
    else:doc.add_heading(text,level=level)
    content.append('\n'+text+'\n')
def table(headers,data,pattern=1):
    node=deepcopy(table_patterns[pattern]);body.insert(len(body)-1,node)
    from docx.table import Table
    t=Table(node,doc._body)
    # Reuse the original table header/body components with the requested rows.
    for r in list(t._tbl.tr_lst)[1:]:t._tbl.remove(r)
    while len(t.columns)<len(headers):t.add_column(Inches(.7))
    if len(t.columns)!=len(headers):
        body.remove(node);t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for i,x in enumerate(headers):t.cell(0,i).text=str(x)
    for data_row in data:
        cells=t.add_row().cells
        for i,x in enumerate(data_row):cells[i].text=str(x)
    for j,row in enumerate(t.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(3)
                for run in paragraph.runs:run.font.name='Arial';run.font.size=Pt(9)
            tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                tag=OxmlElement('w:'+edge);tag.set(qn('w:val'),'single');tag.set(qn('w:sz'),'4');tag.set(qn('w:color'),'D9D9D9');borders.append(tag)
            tcPr.append(borders)
        if j==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    available=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    n=len(headers)
    weights=([.26,.74] if n==2 else [.33,.22,.45] if n==3 else [.31,.23,.23,.23] if n==4 else [.34,.16,.16,.17,.17] if n==5 else [.20,.16,.16,.16,.16,.16])
    if n==6 and headers[0]=='Ref.':weights=[.07,.17,.23,.15,.17,.21]
    t.autofit=False
    for column,weight in zip(t.columns,weights):column.width=int(available*weight)
    for row in t.rows:
        for cell,weight in zip(row.cells,weights):cell.width=int(available*weight)
    content.append('| '+' | '.join(headers)+' |\n'+'\n'.join('| '+' | '.join(str(x) for x in row)+' |' for row in data))
    return t
def figure(name,caption,width=6.0):
    path=assets/name
    if not path.exists():raise FileNotFoundError(path)
    p0=doc.add_paragraph();p0.paragraph_format.keep_with_next=True;p0.add_run().add_picture(str(path),width=Inches(width))
    p(caption);content.append(f'Figure file: {name}')

title='Binary sEMG classification using a temporal proxy of quadriceps fatigue'
p0=p(title,'Title')
for run in p0.runs:run.font.color.rgb=RGBColor(0,0,0)
for element in [p0._p.get_or_add_pPr(),doc.styles['Title'].element.get_or_add_pPr()]:
    for border in list(element.findall(qn('w:pBdr'))):element.remove(border)
p('Reanalysis of the supplied recordings using explicit weak labels and held-out recording groups')
identity={}
if args.previous_report:
    previous=Document(args.previous_report)
    for t in previous.tables:
        for row in t.rows:
            if len(row.cells)>=2 and row.cells[0].text.strip() in ['Group number','Members (name - student ID - role)']:
                identity[row.cells[0].text.strip()]=row.cells[1].text.strip()
commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=repo,text=True).strip()
h('Group identification')
table(['Field','Value'],[
 ['Group number',identity.get('Group number','Not attested in the supplied identity table')],
 ['Report title',title],
 ['Members',identity.get('Members (name - student ID - role)','See original course roster; new individual contributions require attestation')],
 ['Code repository URL','https://github.com/Mateo197802/emg-fatigue-classification'],
 ['Commit hash of this generated version',commit],['Submission date','28 October 2026']])
h('Abstract')
abstract=(f"Surface electromyography can capture changes associated with muscular fatigue, but a classifier must be evaluated against an explicitly defined reference. "
 f"This study implements a binary early-versus-late temporal approximation of acute quadriceps fatigue in the supplied knee-extension recordings. "
 f"The primary archive contains 140 CSV files and 4,256,257 samples grouped into 15 recording-derived subject codes; 138 R1–R5 recordings enter formal processing and two R0 controls remain separate. "
 f"A gap-preserving 20–400 Hz filter, candidate-activity mask, and 250 ms windows with 50% overlap produce {processing['windows_labeled']:,} labeled windows. "
 f"Complete first-third and last-third windows receive the two labels, while transition, invalid-quality and undetected-activity windows are excluded. "
 f"Amplitude, spectral and wavelet features are compared using majority and RMS-plus-MDF logistic baselines, regularized logistic regression, radial-basis SVM and random forest. "
 f"Outer leave-one-recording-group-out validation and grouped inner tuning yield macro-F1 {ms('nested_selector','f1_macro')} for the selection procedure, compared with {ms('logistic_rms_mdf','f1_macro')} for the simple logistic baseline. "
 f"Prespecified window, activity, boundary and notch sensitivities and feature-family ablations quantify dependence on methodological choices. "
 f"Models and preprocessing are serialized and checked after recarga; every table is linked to machine-generated results. "
 f"The central limitation is that no synchronized independent fatigue reference or independently reconciled person identity is available. "
 f"These results therefore demonstrate classification of an operational temporal proxy, rather than confirmed physiological fatigue or generalization to verified new people.")
if not 200<=len(abstract.split())<=250:raise RuntimeError(f'abstract length {len(abstract.split())}')
p(abstract)
h('Keywords');p('surface electromyography; rectus femoris; muscular fatigue proxy; grouped validation; signal processing; reproducibility')
h('1.  Fatigue: definition and state of the art')
h('1.1  Which fatigue are you detecting?',2)
p('The intended construct is acute fatigue during unilateral knee extension, but the available target is a binary elapsed-time approximation. Class 0 denotes eligible early windows; class 1 denotes eligible late windows. Neither central nor peripheral fatigue is identified independently, and no force-loss threshold or clinical fatigue diagnosis is claimed. The assumed exercise-to-failure context belongs to the prescribed protocol and cannot be confirmed per file from the CSV alone. The classification estimates agreement with this weak operational label. It is conditional on valid candidate-active windows in the extremes of a complete recording, not on every window during exercise.')
h('1.2  Physiological basis',2)
p('Under comparable contraction conditions, changes in motor-unit behavior and membrane propagation can shift the sEMG spectrum. Dynamic knee-extension work with synchronized torque, joint angle and MVC offers a materially stronger basis for interpreting spectral fatigue indices than elapsed file position alone [1]. Time-frequency estimation was developed to address the changing signal properties of cyclic contractions [2]. These observations motivate spectral and wavelet feature families without proving that an observed change in the supplied signal is caused by fatigue.')
p('RMS and MAV describe amplitude, which also depends on effort, recruitment, electrode geometry and movement. A decline in MDF or MNF is not specific to fatigue when contraction phase, force, sensor mode and acquisition transfer functions are unknown. Filtering can remove or attenuate components of biological and contaminating signals together; controlled filter studies [3] support explicit processing choices rather than a claim that the filter which maximizes model accuracy is physiologically best. Frequency-analysis recommendations [10] further motivate stated windows, tapering and bandwidth.')
h('1.3  How fatigue has been detected from EMG - review',2)
p('The comparison below uses ten references, including seven primary research articles. DOI metadata were checked; design statements are limited to the source information actually retrieved. Performance values from different muscles, sensors, labels or splits are not imported into the local results. A paper is not classified as subject-independent merely because it uses cross-validation. In particular, the sample-random 10-fold description in [6] is not equivalent to the present grouped design.')
table(['Ref.','Muscle / task','Reference / label','Features','Model','Performance interpretation'],[
 ['[1]','Rectus femoris; dynamic knee extension','Torque, joint angle and pre/post MVC','Spectral indices','No local classifier imported','Different measured reference'],
 ['[2]','Cyclic dynamic contractions','Time-dependent spectral change','Time-frequency','Spectral analysis','No classifier score imported'],
 ['[3]','Controlled sEMG artifact experiment','Noise/filter conditions','Signal spectra','Filter comparison','Not a fatigue-label classifier'],
 ['[4]','sEMG fatigue task','Detailed label not recovered here','Learned feature fusion','CNN/BLSTM attention','Publisher states cross-subject objective; no number imported'],
 ['[5]','sEMG fatigue task','Detailed label/split not recovered here','Wavelet packet','LSTM','No unsupported score imported'],
 ['[6]','Biceps; sEMG+ECG','Self-reported relaxed/transition/fatigue every 120 s','Learned sequences','TCN and evidence fusion','Random sample-fold description; not grouped-equivalent'],
 ['[7]','Gastrocnemius; cycling','Anaerobic-threshold/time phases','sEMG+ACC fusion','CNN attention','Different physiology and modalities'],
 ['[8]','ML science methodology','Leakage taxonomy','Not a fatigue feature study','Methodological analysis','Supports clean separation'],
 ['[9]','Small-sample prediction','Simulation/uncertainty','Not a fatigue label study','Cross-validation analysis','Supports reporting variability'],
 ['[10]','sEMG frequency tutorial','Measurement/analysis choices','Spectral estimation','Tutorial','No classifier score imported']],pattern=2)
h('1.4  Gap and contribution',2)
p('The missing local link is an independent fatigue reference. The contribution is an executable binary pipeline with declared weak supervision, physiologically motivated feature families, grouped nested validation, required baselines, full OOF audit, and actual recargable inference. The historical three-class file-position scores and upper-limb Borg/Cerqueira experiments are retained only in the original archive. They neither train nor validate the new local target.')
h('2.  Dataset and acquisition')
h('2.1  Protocol as executed',2)
p('The supplied guide prescribes bilateral rectus femoris MyoWare 2.0 RAW acquisition at at least 1000 Hz per channel; 20 mm electrode separation; a single preferred-leg 1RM; the same resulting 70% reference load for both legs; 2 s concentric and 2 s eccentric movements to technical failure; Borg CR-10 every five repetitions; ten minutes between legs; resting baseline and pre/post MVC. The same absolute load is not necessarily 70% of each leg’s own 1RM. These are intended conditions. The stored files verify only timestamped signed count values and filenames carrying side/series codes. No synchronized force, video, Borg, electrode photo, load, cadence or failure reason is included. The export column Voltage_mV is ignored because its arithmetic conversion does not verify physical calibration.')
h('2.2  Participants',2)
table(['Variable','Value','Evidence'],[['Analyzed recording groups',15,'Folder-derived; person identity unverified'],['Source recordings',140,'138 formal; two controls'],['Sex, age, mass, 1RM','Not available','Not invented or used as predictors'],['Verified repetitions to failure','Not available','No mechanical markers'],['Additional RAR',20,'Two source codes, kept separate']],pattern=3)
h('2.3  Exclusions',2)
p('Rules were frozen before model evaluation: R0 controls excluded from formal learning; nonfinite values or nonmonotonic timestamps rejected; discontinuities split runs without interpolation; excessive relative timestamp jitter and insufficient rate exclude runs; 100 ms filter edges excluded; windows must be complete and satisfy candidate activity occupancy of at least 75%. Transition and boundary-crossing windows are not a third class. Zero-energy or otherwise undefined features are not filled with zero. No recording is removed because of its classification accuracy.')
totals={k:sum(int(r[k]) for r in coverage) for k in ['raw_samples','qc_valid_samples','candidate_active_samples','candidate_windows','qc_valid_windows','active_windows','feature_valid_windows','labeled_windows']}
table(['Denominator','Count'],[[k.replace('_',' '),f'{v:,}'] for k,v in totals.items()])
h('2.4  Signal quality assessment',2)
p('Quality diagnostics preserve the original bytes and flag clock gaps, jitter, constant segments, extreme-value repetition and spectral line prominence. Signed exported counts do not establish calibrated ADC rails. SNR against seated resting noise is not estimable because no identified resting-baseline segment is provided; filter residual is not substituted for true noise. A record can contribute usable segments while windows crossing a timestamp gap remain excluded.')
figure('qc_examples.png','Figure 1. A continuous accepted example and a gap-crossing interval excluded by the timestamp rule. This is a QC illustration selected without model-outcome information.')
h('3.  Labelling strategy')
h('3.1  Chosen strategy and justification',2)
p('For each original file, u=(t−t_start)/(t_end−t_start). Complete eligible windows ending before u=1/3 receive 0; windows starting at u≥2/3 receive 1. Endpoints are the original observed timestamps, before QC. The middle and boundary-crossing windows are excluded. This implements weak_elapsed_thirds_binary_v1. Repetition-based thirds would be closer to the guide’s example but cannot be reconstructed as verified mechanical events from these files. Borg and MVC approaches were considered and are unavailable locally. The assumption of increasing fatigue with elapsed file time is explicitly weaker than either independent approach.')
h('3.2  Class definition and distribution',2)
counts=processing['label_counts'];n=sum(counts.values())
table(['Class','Exact eligible support','Windows','Percent'],[['0 early proxy','Complete first elapsed-time third',counts['0'],f"{100*counts['0']/n:.1f}%"],['1 late proxy','Complete last elapsed-time third',counts['1'],f"{100*counts['1']/n:.1f}%"]],pattern=5)
h('3.3  Circularity check',2)
p('No EMG feature defines label 0 versus 1: labels depend on elapsed recording bounds. RMS enters the candidate-activity eligibility mask, which creates a selection condition, not a fatigue class. This can still bias the analyzed population of windows and is addressed by activity-occupancy sensitivity. Time, file name, recording/group identity, side and series are not model inputs. MDF/MNF threshold-screen outputs from the original project are not labels or predictors in this task. The time oracle diagnostic is tautological by construction and is excluded from deployable models.')
h('3.4  Label reliability',2)
p('No independent local label source exists, so inter-source kappa cannot be estimated. Classifier-versus-proxy kappa is reported separately and is not label reliability. Boundary shifts of ±5% and ±10% recording duration re-evaluate the fixed RMS+MDF baseline. These changes quantify dependence on arbitrary temporal boundaries, not validation against force or subjective fatigue. A one-repetition shift cannot be evaluated without repetition markers.')
h('4.  Signal processing and feature extraction')
h('4.1  Preprocessing chain',2)
table(['Stage','Method','Parameters / limitation'],[['Band pass','Butterworth SOS, bidirectional','20–400 Hz; prototype order 4; 100 ms edge mask'],['Notch','Not primary; prespecified sensitivity','50/60 Hz prominence >10 dB; Q=30'],['Rectification','Separate RMS activity envelope','Unrectified stored-signal branch for PSD'],['Amplitude normalization','Train-fold StandardScaler','No fabricated %MVC'],['Artefact handling','Gap, jitter, constant-run checks','No crossing discontinuities; signed rails uncalibrated'],['Segmentation','Candidate activity, not repetitions','50 ms RMS; 25 ms hop; per-record Otsu'],['Window','Complete contiguous windows','250 ms; 50% overlap; occupancy ≥75%'],['Resampling / detrending','No resampling; mean removal','Nominal fs from median dt']],pattern=6)
h('4.2  Features computed',2)
table(['Feature','Family','Definition','Reason / condition'],[['RMS','Amplitude','sqrt(mean((x−mean(x))²))','Effective amplitude; effort and placement confound'],['MAV','Amplitude','mean(abs(x−mean(x)))','Amplitude distribution'],['Waveform length/s','Amplitude','sum(abs(diff(x)))/duration','Local variation'],['MDF','Frequency','50% cumulative Hann PSD in band','Spectral redistribution, not specific fatigue'],['MNF','Frequency','sum(fP)/sum(P)','Spectral mean'],['Low/high ratio','Frequency','[20,60)/[60,400] Hz power','Relative spectrum shape'],['A3,D3,D2,D1','Time-frequency','db4 level 3 relative energies','Multiscale variation under nonstationarity']],pattern=7)
h('4.3  Feature behaviour across the set',2)
figure('trajectories.png','Figure 2. Candidate-active feature trajectories by recording fraction. First reduce to recording-bin medians, then group-bin medians; bands show across-group SD. They do not establish force loss.')
h('4.4  Bilateral analysis',2)
p('Left and right recordings remain in one group for every split. Feature summaries pair side within group and series where both are available. There is no evidence that rows from separately exported sides are synchronous, and no paired-sample concatenation is used. Side differences are electrical summaries, not proof of dominant leg or unequal fatigue rates.')
h('4.5  Feature selection or reduction',2)
p('All-feature models use the ten prespecified features; the required simple baseline uses exactly RMS and MDF. No data-driven feature subset is selected on all groups. Family ablations remove amplitude, spectral or wavelet features and repeat grouped inner logistic tuning. Scaling is part of the fitted Pipeline and sees training rows only [8].')
h('5.  Machine learning methodology')
h('5.1  Models evaluated',2)
table(['Model','Rationale','Search','Selection'],[['Majority','Trivial mandatory floor','Training majority','No test-derived prior'],['Logistic RMS+MDF','Mandatory simple reference','C=0.1/1/10','Grouped inner macro-F1'],['Logistic all','Regularized linear reference','C=0.1/1/10','Grouped inner macro-F1'],['SVM RBF','Nonlinear interaction','C=0.1/1/10; gamma=scale/0.01','Grouped inner macro-F1'],['Random forest','Nonlinear feature partition','300 trees; depth=3/6/None; leaf=2/5','Grouped inner macro-F1'],['Nested selector','Evaluate family-selection procedure','Across the same candidate grid','Inner performance only']],pattern=9)
h('5.2  Validation scheme',2)
p('Fifteen outer folds each hold out one recording group completely. Four inner grouped folds select configurations using average per-validation-group macro-F1. Both sides and all series of each group remain together. Test-group observations never fit the scaler, model, hyperparameters, threshold or feature subset. The candidate-activity and deterministic filtering transforms are performed independently per recording under the frozen offline protocol. Folder identity is not independently reconciled, so the unseen-person interpretation remains conditional. OOF predictions cover every eligible window exactly once per model. Small sample uncertainty remains substantial [9].')
figure('split_scheme.png','Figure 3. Outer group holdout and grouped inner selection. Learned preprocessing is fitted only on the training branch.')
h('5.3  Class imbalance',2)
p(f"The analyzed weak-label distribution is {counts['0']} early and {counts['1']} late windows. Logistic, SVM and random-forest candidates use balanced class weights computed within fitting. Majority class is determined from training labels only. No synthetic resampling combines subjects or creates test observations. Group-level evaluation means are not weighted by how many windows a person contributed.")
h('5.4  Reproducibility settings',2)
env=processing['environment'];total_runtime=sum(float(r['fit_predict_seconds']) for r in results if r['model']=='nested_selector')
table(['Item','Value'],[['Seeds','20261028'],['Python and numeric libraries',json.dumps(env,sort_keys=True)],['Hardware',platform.processor()+'; CPU sklearn execution'],['Nested evaluation wall time',f'{total_runtime:.1f} s summed across outer fold timers'],['Regenerate assets','python -m emg_fatigue report --config configs/study.yaml'],['Processing configuration SHA256',processing['config_sha256']]],pattern=11)
timers=[float(r['fit_predict_seconds']) for r in results if r['model']=='nested_selector']
if max(timers)>3*sorted(timers)[len(timers)//2]:p(f'The measured wall timers include an exceptional fold interval of {max(timers):.1f} s. No CPU-only time or steady-state training benchmark is inferred from that interval; it remains included in the audit rather than being removed.')
h('6.  Results')
h('6.1  Main comparison table',2)
models=['majority','logistic_rms_mdf','logistic_all','svm_rbf','random_forest','nested_selector']
table(['Model','Accuracy','Precision macro','Recall macro','F1 macro','ROC AUC'],[[m]+[ms(m,k) for k in ['accuracy','precision_macro','recall_macro','f1_macro','roc_auc']] for m in models],pattern=12)
p(f"The prespecified nested procedure achieved macro-F1 {ms('nested_selector','f1_macro')} (group-bootstrap 95% CI {ci('nested_selector','f1_macro')}). Average precision was {ms('nested_selector','average_precision')}; trapezoidal PR-AUC was {ms('nested_selector','pr_auc_trapezoidal')}. The RMS+MDF baseline achieved {ms('logistic_rms_mdf','f1_macro')}. Family point estimates do not establish a unique superior model, and choosing a winner from these outer estimates would add post-selection optimism.")
for r in paired:p(f"Paired contrast {r['comparison']}: mean macro-F1 difference {float(r['mean_difference']):+.3f}; 95% group-bootstrap interval {float(r['ci95_lower']):+.3f} to {float(r['ci95_upper']):+.3f}, {r['paired_groups']} matched groups.")
h('6.2  Per-class performance of the best model',2)
p('Here “best model” means the prespecified nested selection procedure, not the family with the largest retrospectively observed point estimate. Class metrics are mean ± SD across held-out groups; support below is the pooled OOF count.')
table(['Class','Precision','Recall','F1','Support'],[[f'{i} '+('early proxy' if i==0 else 'late proxy'),ms('nested_selector',f'precision_class{i}'),ms('nested_selector',f'recall_class{i}'),ms('nested_selector',f'f1_class{i}'),pooled[f'support_class{i}']] for i in [0,1]],pattern=13)
h('6.3  Confusion matrix',2)
figure('confusion.png','Figure 4. Pooled OOF confusion counts and row-normalized proportions. Window pooling does not replace across-group uncertainty.')
h('6.4  Per-subject performance',2)
figure('held_out_groups.png','Figure 5. One point per held-out recording group; identity follows source folders and is unverified independently. The figure remains in the private report.')
h('6.5  Feature importance or ablation',2)
figure('ablation.png','Figure 6. Logistic feature-family ablations with grouped inner tuning. Lower performance after removal indicates predictive contribution for the proxy, not a causal physiological mechanism.')
table(['Subset','Macro F1 mean','SD'],[[r['ablation'],f"{float(r['macro_f1_mean']):.3f}",f"{float(r['macro_f1_sd']):.3f}"] for r in abl])
h('6.6  Bilateral results',2)
bil=rows(assets/'bilateral_summary.csv')
table(['Feature','Paired group-series','Mean L−R','Median L−R'],[[r['feature'],r['paired_group_series'],f"{float(r['mean_left_minus_right']):.3f}",f"{float(r['median_left_minus_right']):.3f}"] for r in bil])
p('These comparisons summarize candidate-active electrical features in available matched side/series combinations. A group can contribute several series, so the pair count is not an independent-person sample size. Held-out-group side-specific model metrics are provided privately. No dominant-leg, same-relative-intensity or unequal-fatigue-rate conclusion is justified.')
h('7.  Discussion')
h('7.1  Interpretation',2)
p('The implementation now answers a declared binary operational task with actual inference, but its biological interpretation remains limited by supervision. The exact-time oracle reaches perfect agreement by definition and demonstrates why a high temporal-label score would not itself validate fatigue. The nested model score does not clearly exceed the simple RMS+MDF baseline. The final full-cohort artifact is selected using internal grouped validation, and is not an additional generalization experiment. It should be used as an auditable research prototype for this weak target.')
h('7.2  Comparison with the literature',2)
p('The torque- and MVC-informed knee-extension experiment in [1] addresses a stronger physiological reference than these files. Dynamic spectral estimation [2] motivates attention to nonstationarity and contraction conditions. Fusion methods [4–7] use other architectures, modalities or targets, so their published accuracy cannot be ranked directly against the current grouped macro-F1. The random-sample fold description of [6] further illustrates why a cross-validation label does not guarantee subject separation. We adopt leakage controls [8] and expose small-cohort variability [9] rather than importing a literature score as local evidence.')
h('7.3  Failure analysis',2)
selected=[r for r in results if r['model']=='nested_selector'];worst=min(selected,key=lambda r:float(r['f1_macro']));best=max(selected,key=lambda r:float(r['f1_macro']))
p(f"Held-out-group macro-F1 ranged from {float(worst['f1_macro']):.3f} to {float(best['f1_macro']):.3f}. The least accurate fold contained {worst['n_test']} eligible windows, with confusion counts TN={worst['cm_0_0']}, FP={worst['cm_0_1']}, FN={worst['cm_1_0']}, TP={worst['cm_1_1']}. These are measured error counts, not a diagnosis of its cause. Misclassified windows are preserved in private OOF predictions for review. Without video, force or acquisition metadata, attributing the errors to true fatigue, technique or electrode displacement would be speculation.")
h('7.4  Limitations',2)
p('Labels are approximate and cannot establish an absence of fatigue in the first third. Group identity is inferred from filenames; person-level independence remains conditional. The cohort is small, signals overlap, phase and load are not verified, and candidate activity may include artefact. Transition and low-activity windows are outside the evaluated target, so reported performance is conditional on substantial selection. Physical channel mode and calibration, synchronized force/MVC, Borg and mechanical repetition markers are unavailable. Every transform is offline; zero-phase filtering, per-recording activity thresholds and knowledge of full recording bounds prevent a real-time onset claim. The separate RAR is not independently identified or untouched external validation. Some literature label/split details remain unrecovered and are explicitly marked rather than assumed.')
h('7.5  What you would do differently',2)
p('Before new acquisition, obtain the existing course coordinator’s acquisition metadata and synchronized references if available under the original consent and protocol. Reconcile person/session keys and sensor mode, then define the reference construct before looking at model results. If that information cannot be recovered, retain the present weak-target scope and avoid renaming elapsed-time classification as physiological validation. Any new acquisition would require its own approved course protocol and cannot silently replace the shared cohort.')
h('8.  Conclusions')
p(f"1. An executable binary pipeline produced {processing['windows_labeled']:,} eligible weak-labeled windows from 138 formal recordings, with all quality and selection denominators retained.")
p(f"2. The nested selection procedure achieved macro-F1 {ms('nested_selector','f1_macro')}; its observed mean did not clearly improve on the {ms('logistic_rms_mdf','f1_macro')} RMS+MDF baseline.")
p('3. Signal processing, grouped OOF evaluation, complete model recarga and fragment reassembly are auditable software capabilities; they do not substitute for a fatigue reference.')
p('4. The available evidence supports early/late proxy classification. Independent physiological fatigue detection, calibrated force loss and verified unseen-person generalization remain unestablished.')
h('9.  Contribution of each member')
p('The original roster is retained above when available from the supplied report. The new repository records assisted implementation under the authenticated account used for execution. It does not prove balanced student effort. Individual roles and effort percentages for the rebuilt work require the members’ own attestation; no authors, roles, percentages or retrospective commits have been fabricated. This administrative requirement remains outstanding for course submission.')
h('10.  References')
for i,r in enumerate(refs,1):p(f"[{i}] {r['authors']}, “{r['title']},” {r['journal']}, {r['year']}. doi: {r['doi']}.")
h('Appendix A - Additional figures')
figure('sensitivity.png','Prespecified sensitivity analyses, fixed C=1 RMS+MDF logistic reference. Different retained-window populations are stated explicitly.')
table(['Variation','Windows','Groups','Macro F1','SD'],[[r['analysis'],r.get('n_windows','NA'),r.get('groups','NA'),f"{float(r['macro_f1_mean']):.3f}" if r.get('macro_f1_mean') else 'NA',f"{float(r['macro_f1_sd']):.3f}" if r.get('macro_f1_sd') else 'NA'] for r in sens])
h('Appendix B - Full feature list')
p('rms; mav; waveform_length_per_s; mdf_hz; mnf_hz; power_ratio_low_high; wavelet_A3; wavelet_D3; wavelet_D2; wavelet_D1. Metadata are excluded by an explicit column transformer. The all-feature schema is identical in fitting and prediction.')
h('Appendix C - Repository map')
p('io/contracts: source and configuration contracts. qc/preprocessing/activity/labels/features: data transformations. validation/evaluation: nested grouped OOF and metrics. inference: trusted local bundle and prediction. sensitivities: frozen variations and family ablations. reporting: numerical assets and sources. tests: software contracts and regressions. Raw signals, models, OOF predictions and this full report remain private; public code and deidentified fold metrics are separate.')
for section in doc.sections:
    for paragraph in section.footer.paragraphs:
        for run in paragraph.runs:
            if 'Report Template' in run.text:run.text=run.text.replace('Report Template','Binary sEMG reanalysis')
final=manuscript/'EMG_Binary_Report.docx';doc.save(final)
(manuscript/'EMG_Binary_Report.md').write_text('\n\n'.join(content)+'\n',encoding='utf-8')
if template.read_bytes()!=template_bytes:raise RuntimeError('reference template changed')
write={'template_sha256':hashlib.sha256(template_bytes).hexdigest(),'template_geometry':geometry,'output_sha256':hashlib.sha256(final.read_bytes()).hexdigest(),'code_commit':commit,'abstract_word_count':len(abstract.split()),'top_level_sections_preserved':list(range(1,11)),'body_edit':'all editable template body slots filled; blank placeholder spacing removed; source styles/header/footer/section retained','unresolved':['independent fatigue reference','person identity reconciliation','physical calibration/mode','individual student contribution attestation','selected literature label/split details explicitly marked']}
(manuscript/'manuscript_receipt.json').write_text(json.dumps(write,indent=2)+'\n',encoding='utf-8')
print(final)
